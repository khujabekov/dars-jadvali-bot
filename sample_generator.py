"""
jadval_namuna.xlsx faylini chiroyli dizayn bilan yaratuvchi skript.
"""
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DATA_DIR = Path(__file__).parent / "data"
TEMPLATE_FILE = DATA_DIR / "jadval_namuna.xlsx"

def create_sample_excel(target_path: Path = TEMPLATE_FILE) -> Path:
    target_path.parent.mkdir(parents=True, exist_ok=True)
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dars Jadvali"
    
    # Sarlavhalar
    headers = ["Kun", "Vaqt", "Fan", "O'qituvchi", "Xona", "Dars turi"]
    ws.append(headers)
    
    # Namunaviy darslar
    sample_data = [
        # Dushanba
        ["Dushanba", "08:30 - 09:50", "Oliy matematika", "Dots. Karimov A.", "305-xona", "Ma'ruza"],
        ["Dushanba", "10:00 - 11:20", "Dasturlash asoslari (Python)", "Ass. Alimov B.", "210-kompyuter", "Amaliyot"],
        ["Dushanba", "11:30 - 12:50", "Ingliz tili", "Katta o'qit. Karimova D.", "104-xona", "Amaliyot"],
        
        # Seshanba
        ["Seshanba", "08:30 - 09:50", "Fizika", "Prof. Rahimov S.", "B-korpus 201", "Ma'ruza"],
        ["Seshanba", "10:00 - 11:20", "Ma'lumotlar tuzilmasi", "Dots. Sultonov M.", "212-kompyuter", "Laboratoriya"],
        ["Seshanba", "11:30 - 12:50", "Falsafa", "Ass. Ergashev N.", "401-xona", "Ma'ruza"],
        
        # Chorshanba
        ["Chorshanba", "08:30 - 09:50", "Kiberxavfsizlik asoslari", "Dots. Yusupov T.", "302-xona", "Ma'ruza"],
        ["Chorshanba", "10:00 - 11:20", "Algoritmlar nazariyasi", "Ass. Alimov B.", "210-kompyuter", "Amaliyot"],
        ["Chorshanba", "11:30 - 12:50", "Jismoniy tarbiya", "Katta o'qit. Nazarov X.", "Sport zal", "Amaliyot"],
        
        # Payshanba
        ["Payshanba", "08:30 - 09:50", "Ma'lumotlar bazasi (SQL)", "Dots. Saidov K.", "215-kompyuter", "Ma'ruza"],
        ["Payshanba", "10:00 - 11:20", "Ma'lumotlar bazasi (SQL)", "Ass. Vohidov Z.", "215-kompyuter", "Laboratoriya"],
        ["Payshanba", "11:30 - 12:50", "Oliy matematika", "Dots. Karimov A.", "305-xona", "Amaliyot"],
        
        # Juma
        ["Juma", "08:30 - 09:50", "Web dasturlash", "Ass. Ismoilov D.", "208-kompyuter", "Ma'ruza"],
        ["Juma", "10:00 - 11:20", "Web dasturlash", "Ass. Ismoilov D.", "208-kompyuter", "Amaliyot"],
        
        # Shanba
        ["Shanba", "08:30 - 09:50", "Sun'iy intellektga kirish", "Prof. Qodirov O.", "Akt zali", "Ma'ruza"],
        ["Shanba", "10:00 - 11:20", "Mustaqil ta'lim / To'garak", "Kafedra o'qituvchilari", "Kutubxona", "Mustaqil ish"],
    ]
    
    for row in sample_data:
        ws.append(row)
        
    # Styling (Bezatish)
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=11)
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    # Sarlavhani formatlash
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Qatorlarni formatlash (kunlar bo'yicha orqa fon rangini o'zgartirish)
    day_colors = {
        "Dushanba": "F2F5F9",
        "Seshanba": "FFFFFF",
        "Chorshanba": "F2F5F9",
        "Payshanba": "FFFFFF",
        "Juma": "F2F5F9",
        "Shanba": "FFFFFF"
    }
    
    for row in ws.iter_rows(min_row=2, max_row=len(sample_data)+1, min_col=1, max_col=len(headers)):
        day_val = row[0].value
        row_color = day_colors.get(day_val, "FFFFFF")
        row_fill = PatternFill(start_color=row_color, end_color=row_color, fill_type="solid")
        
        for cell in row:
            cell.font = data_font
            cell.fill = row_fill
            cell.border = thin_border
            if cell.column in [1, 2, 5, 6]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                
    # Ustun kengliklarini avtomatik moslash
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
        
    ws.row_dimensions[1].height = 26
    for r in range(2, len(sample_data) + 2):
        ws.row_dimensions[r].height = 22
        
    wb.save(target_path)
    print(f"Namuna fayli saqlandi: {target_path}")
    return target_path

if __name__ == "__main__":
    create_sample_excel()
