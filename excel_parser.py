import html
from pathlib import Path
from datetime import datetime, date
from typing import Dict, List, Optional, Tuple
import openpyxl
import xlrd

UZBEK_DAYS = {
    0: "Dushanba",
    1: "Seshanba",
    2: "Chorshanba",
    3: "Payshanba",
    4: "Juma",
    5: "Shanba",
    6: "Yakshanba"
}

NUM_EMOJIS = {
    1: "1️⃣", 2: "2️⃣", 3: "3️⃣", 4: "4️⃣", 5: "5️⃣",
    6: "6️⃣", 7: "7️⃣", 8: "8️⃣", 9: "9️⃣", 10: "🔟"
}

def clean_str(val) -> str:
    if val is None:
        return ""
    s = str(val).strip()
    if s.endswith(".0") and s[:-2].isdigit():
        return s[:-2]
    return s

def escape_html(val: any) -> str:
    s = clean_str(val)
    return html.escape(s, quote=False)

def format_room(room_str: any) -> str:
    r = clean_str(room_str)
    if not r:
        return ""
    if r.isdigit():
        return f"{r}-xona"
    return r

def normalize_day(day_str: any) -> Optional[int]:
    """Kun matnini 0-6 oraliqdagi raqamga aylantiradi (0 = Dushanba, 6 = Yakshanba)."""
    if not day_str:
        return None
    
    val = str(day_str).strip().lower()
    if val in ("1", "1-kun"): return 0
    if val in ("2", "2-kun"): return 1
    if val in ("3", "3-kun"): return 2
    if val in ("4", "4-kun"): return 3
    if val in ("5", "5-kun"): return 4
    if val in ("6", "6-kun"): return 5
    if val in ("7", "7-kun"): return 6
    
    if "dush" in val or "mon" in val or "пон" in val:
        return 0
    if "sesh" in val or "tue" in val or "вто" in val:
        return 1
    if "chor" in val or "wed" in val or "сре" in val:
        return 2
    if "pay" in val or "thu" in val or "чет" in val:
        return 3
    if "jum" in val or "fri" in val or "пят" in val:
        return 4
    if "shan" in val or "sat" in val or "суб" in val:
        return 5
    if "yak" in val or "sun" in val or "вос" in val:
        return 6
        
    return None

def parse_schedule_excel(file_path: Path) -> Tuple[Dict[int, List[dict]], Optional[str]]:
    """
    Excel faylni (.xlsx yoki .xls) o'qiydi va kunlar bo'yicha lug'at (0-6) qaytaradi.
    Texnikum / kollej (juftliklar va kichik guruhlarga bo'lingan) hamda oddiy jadval formatlarini qo'llab-quvvatlaydi.
    """
    if not file_path.exists():
        return {}, "Jadval fayli topilmadi."

    ext = file_path.suffix.lower()

    if ext == ".xls":
        try:
            wb = xlrd.open_workbook(str(file_path), formatting_info=True)
            ws = wb.sheet_by_index(0)
            nrows, ncols = ws.nrows, ws.ncols
            raw_grid = [[ws.cell_value(r, c) for c in range(ncols)] for r in range(nrows)]
            grid = [[ws.cell_value(r, c) for c in range(ncols)] for r in range(nrows)]
            for rlo, rhi, clo, chi in ws.merged_cells:
                top_val = ws.cell_value(rlo, clo)
                for r in range(rlo, rhi):
                    for c in range(clo, chi):
                        grid[r][c] = top_val
        except Exception as e:
            return {}, f"XLS faylini o'qishda xatolik: {e}"
    else:
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            ws = wb.active
            if not ws:
                return {}, "Excel faylida sahifa topilmadi."
            nrows, ncols = ws.max_row, ws.max_column
            raw_grid = [[ws.cell(r, c).value for c in range(1, ncols + 1)] for r in range(1, nrows + 1)]
            grid = [[ws.cell(r, c).value for c in range(1, ncols + 1)] for r in range(1, nrows + 1)]
            for m in ws.merged_cells.ranges:
                top_val = ws.cell(m.min_row, m.min_col).value
                for r in range(m.min_row, m.max_row + 1):
                    for c in range(m.min_col, m.max_col + 1):
                        grid[r - 1][c - 1] = top_val
        except Exception as e:
            return {}, f"Excel faylini ochishda xatolik: {e}"

    # Sarlavhalarni aniqlash
    header_row_idx = None
    col_map = {}
    is_college_format = False

    for r_idx in range(min(10, nrows)):
        row_str = [clean_str(raw_grid[r_idx][c]).lower() for c in range(ncols)]
        for c_idx, text in enumerate(row_str):
            if any(k in text for k in ["kun", "day", "hafta"]):
                col_map["day"] = c_idx
            elif any(k in text for k in ["vaqt", "soat", "time"]):
                col_map["time"] = c_idx
            elif any(k in text for k in ["fan", "subject", "lesson"]):
                col_map["subject"] = c_idx
            elif any(k in text for k in ["o'qituvchi", "ustoz", "domla", "teacher", "f.i.sh"]):
                col_map["teacher"] = c_idx
            elif any(k in text for k in ["xona", "auditoriya", "room"]):
                col_map["room"] = c_idx
            elif any(k in text for k in ["juftlik", "para"]):
                col_map["pair"] = c_idx
                is_college_format = True
            elif any(k in text for k in ["turi", "type"]):
                col_map["type"] = c_idx

        if "day" in col_map and "subject" in col_map:
            header_row_idx = r_idx
            break

    if header_row_idx is None:
        return {}, "Excel faylida 'Kun' va 'Fan' ustunlari topilmadi."

    schedule: Dict[int, List[dict]] = {i: [] for i in range(7)}

    if is_college_format or "pair" in col_map:
        day_col = col_map.get("day", 0)
        pair_col = col_map.get("pair", 1)
        time_col = col_map.get("time", 2)
        sub_col = col_map.get("subject", 3)
        room_col = col_map.get("room", 4)
        teacher_col = col_map.get("teacher", 5)

        r = header_row_idx + 1
        while r < nrows:
            day_val = clean_str(grid[r][day_col])
            parsed_day = normalize_day(day_val)
            if parsed_day is None:
                r += 1
                continue

            pair_val = clean_str(grid[r][pair_col])
            time_val = clean_str(grid[r][time_col])

            r1 = r
            r2 = r + 1 if r + 1 < nrows and clean_str(grid[r + 1][pair_col]) == pair_val else r
            r = r2 + 1

            sub1 = clean_str(raw_grid[r1][sub_col]) or clean_str(grid[r1][sub_col])
            room1 = format_room(raw_grid[r1][room_col] if raw_grid[r1][room_col] is not None else grid[r1][room_col])
            teacher1 = clean_str(raw_grid[r1][teacher_col]) or clean_str(grid[r1][teacher_col])

            sub2 = clean_str(raw_grid[r2][sub_col]) or clean_str(grid[r2][sub_col])
            room2 = format_room(raw_grid[r2][room_col] if raw_grid[r2][room_col] is not None else grid[r2][room_col])
            teacher2 = clean_str(raw_grid[r2][teacher_col]) or clean_str(grid[r2][teacher_col])

            if not sub1 and not sub2:
                continue

            # Agar ikkala kichik guruh bitta dars bo'lsa
            if (sub1 == sub2 or not sub2) and (room1 == room2 or not room2) and (teacher1 == teacher2 or not teacher2):
                schedule[parsed_day].append({
                    "time": time_val,
                    "subject": sub1,
                    "room": room1,
                    "teacher": teacher1,
                    "type": ""
                })
            else:
                schedule[parsed_day].append({
                    "time": time_val,
                    "subject": sub1 if sub1 == sub2 else f"{sub1} / {sub2}",
                    "room": f"{room1} / {room2}" if room1 != room2 else room1,
                    "teacher": f"{teacher1} / {teacher2}" if teacher1 != teacher2 else teacher1,
                    "subgroups": [
                        {"subject": sub1, "room": room1, "teacher": teacher1},
                        {"subject": sub2, "room": room2, "teacher": teacher2}
                    ]
                })
    else:
        day_col = col_map["day"]
        sub_col = col_map["subject"]
        time_col = col_map.get("time")
        room_col = col_map.get("room")
        teacher_col = col_map.get("teacher")
        type_col = col_map.get("type")

        current_day = None
        for r in range(header_row_idx + 1, nrows):
            day_cell = clean_str(grid[r][day_col])
            sub_cell = clean_str(grid[r][sub_col])

            if day_cell:
                d = normalize_day(day_cell)
                if d is not None:
                    current_day = d

            if not sub_cell or current_day is None:
                continue

            time_val = clean_str(grid[r][time_col]) if time_col is not None else ""
            room_val = format_room(grid[r][room_col]) if room_col is not None else ""
            teacher_val = clean_str(grid[r][teacher_col]) if teacher_col is not None else ""
            type_val = clean_str(grid[r][type_col]) if type_col is not None else ""

            schedule[current_day].append({
                "time": time_val,
                "subject": sub_cell,
                "room": room_val,
                "teacher": teacher_val,
                "type": type_val
            })

    return schedule, None

def format_day_schedule(
    day_num: int, 
    target_date: Optional[date] = None, 
    schedule: Optional[Dict[int, List[dict]]] = None, 
    file_path: Optional[Path] = None,
    header_title: Optional[str] = None
) -> str:
    """Kunlik dars jadvalini so'ralgan formatda (sana | KUN, vaqt va fan) qaytaradi."""
    if schedule is None:
        if file_path is None:
            from config import SCHEDULE_FILE
            file_path = SCHEDULE_FILE
        schedule, error = parse_schedule_excel(file_path)
        if error:
            return f"❌ <b>Xatolik:</b> {escape_html(error)}"
            
    day_name = UZBEK_DAYS.get(day_num, "Noma'lum kun").upper()
    date_str = target_date.strftime("%d.%m.%Y") if target_date else ""
    header = f"<b>{date_str} | {day_name}</b>" if date_str else f"<b>{day_name}</b>"
    
    lessons = schedule.get(day_num, [])
    
    if not lessons:
        if day_num == 6:  # Yakshanba
            empty_msg = "Bugun dam olish kuni."
        else:
            empty_msg = "Bugun dars yo'q."
        return f"{header}\n\n{empty_msg}"
            
    items = []
    for item in lessons:
        time_raw = clean_str(item.get('time', ''))
        start_time = time_raw.split('-')[0].strip() if time_raw else ""
        time_part = f"<b>{escape_html(start_time)}</b>   " if start_time else ""
        
        if "subgroups" in item:
            sg = item["subgroups"]
            s1 = sg[0]
            s2 = sg[1]
            sub1_name = escape_html(s1.get('subject', ''))
            sub2_name = escape_html(s2.get('subject', ''))
            
            if s1.get('subject') == s2.get('subject'):
                items.append(f"{time_part}{sub1_name}")
            else:
                items.append(f"{time_part}{sub1_name} / {sub2_name}")
        else:
            subject_name = escape_html(item.get('subject', ''))
            items.append(f"{time_part}{subject_name}")
        
    return f"{header}\n\n" + "\n".join(items)

def format_week_schedule(file_path: Optional[Path] = None) -> str:
    """Haftalik to'liq dars jadvalini ixcham ko'rinishda qaytaradi."""
    if file_path is None:
        from config import SCHEDULE_FILE
        file_path = SCHEDULE_FILE
        
    schedule, error = parse_schedule_excel(file_path)
    if error:
        return f"❌ <b>Xatolik:</b> {escape_html(error)}"
        
    total_lessons = sum(len(l) for l in schedule.values())
    if total_lessons == 0:
        return "⚠️ Dars jadvali kiritilmagan."
        
    lines = ["<b>HAFTALIK DARS JADVALI</b>\n"]
    
    for day_num in range(7):
        lessons = schedule.get(day_num, [])
        day_name = UZBEK_DAYS[day_num].upper()
        
        if not lessons:
            continue
            
        lines.append(f"<b>{day_name}:</b>")
        for item in lessons:
            time_raw = clean_str(item.get('time', ''))
            start_time = time_raw.split('-')[0].strip() if time_raw else ""
            time_part = f"<b>{escape_html(start_time)}</b>   " if start_time else ""
            
            if "subgroups" in item:
                sg = item["subgroups"]
                s1 = sg[0]
                s2 = sg[1]
                sub1_name = escape_html(s1.get('subject', ''))
                sub2_name = escape_html(s2.get('subject', ''))
                if s1.get('subject') == s2.get('subject'):
                    lines.append(f"{time_part}{sub1_name}")
                else:
                    lines.append(f"{time_part}{sub1_name} / {sub2_name}")
            else:
                subject_name = escape_html(item.get('subject', ''))
                lines.append(f"{time_part}{subject_name}")
            
        lines.append("")  # Bo'sh qator
        
    return "\n".join(lines).strip()

def convert_xls_to_xlsx(src_path: Path, dst_path: Path):
    """Eski .xls faylni .xlsx formatiga o'tkazadi va birlashtirilgan kataklarni saqlaydi."""
    wb_xls = xlrd.open_workbook(str(src_path), formatting_info=True)
    sheet_xls = wb_xls.sheet_by_index(0)
    
    wb_xlsx = openpyxl.Workbook()
    ws_xlsx = wb_xlsx.active
    ws_xlsx.title = sheet_xls.name
    
    for r in range(sheet_xls.nrows):
        row_vals = [sheet_xls.cell_value(r, c) for c in range(sheet_xls.ncols)]
        ws_xlsx.append(row_vals)
        
    for rlo, rhi, clo, chi in sheet_xls.merged_cells:
        ws_xlsx.merge_cells(
            start_row=rlo + 1,
            end_row=rhi,
            start_column=clo + 1,
            end_column=chi
        )
        
    wb_xlsx.save(str(dst_path))

